"""Build a screening step's dataset from its LLM judges and manual review.

A repo passes when:
  - all judges gave it the same label and that label is the wanted one, or
  - the judges disagree and its "Final Decision" in manual_review.xlsx is the
    wanted label. The manual review is only consulted when the judges disagree.

Writes an .xlsx (summary + passing repos) and a JSON in the same format as
dataset/all_repos.json containing only the passing repos. --step picks the defaults
(judges, label, input, manual review and output) for S1 or S2.

Usage:
    python merge_judges.py                 # S1: app vs framework
    python merge_judges.py --step s2       # S2: app cleanup, on the S1 output
    python merge_judges.py --label framework -o ../../../dataset/s1_frameworks
"""
import argparse
import glob
import json
import os
from collections import Counter
from datetime import datetime
from itertools import combinations

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
SCREENING_DIR = os.path.join(PROJECT_ROOT, 'crawler', 'output', 'llm_screening')
DATASET_DIR = os.path.join(PROJECT_ROOT, 'dataset')
# Per-step defaults; judges=None uses every file in the step's judges/ folder
STEPS = {
    's1': dict(dir='s1_app_vs_framework', label='app', repos='all_repos.json', output='s1_apps',
               judges=['deepseek=deepseek-r1.json', 'llama=llama3.1.json', 'gpt=gpt-5.4-mini.json']),
    's2': dict(dir='s2_app_cleanup', label='Agentic Applications', repos='s1_apps.json',
               output='s2_agentic_apps', judges=None),
}
UNDECIDED = {'', 'error', 'needs_review', 'needs review'}  # judge answers that are not a label
STAR_BUCKETS = [(0, '< 10'), (10, '10 - 99'), (100, '100 - 999'), (1000, '1k - 9.9k'), (10000, '10k+')]
HEADER_FILL = PatternFill(start_color='DCE8F7', end_color='DCE8F7', fill_type='solid')


def resolve_path(path, judges_dir):
    return path if os.path.exists(path) else os.path.join(judges_dir, path)


def load_judge(spec, judges_dir):
    """Parse 'name=path' (or just 'path') and return (name, {url: label})."""
    name, _, path = spec.rpartition('=')
    with open(resolve_path(path, judges_dir)) as f:
        data = json.load(f)
    name = name or data.get('model', os.path.basename(path)).split(':')[0].replace('gpt-gpt-', 'gpt-')
    return name, {r['url']: 'error' if r.get('error') or not r.get('classification') else r['classification']
                  for r in data['results']}


def is_formula(value):
    return str(value).startswith('=')


def load_manual(path):
    """Return {url: decision} for every row of manual_review.xlsx with a typed-in Final Decision,
    or {} when the step has no manual review yet."""
    if not os.path.exists(path):
        return {}
    rows = load_workbook(path)['Classifications'].iter_rows(values_only=True)
    col = {str(h).strip(): i for i, h in enumerate(next(rows)) if h}
    return {r[col['Repository URL']]: str(r[col['Final Decision']]).strip().lower()
            for r in rows
            if str(r[col['Repository URL']]).startswith('http')
            and r[col['Final Decision']] and not is_formula(r[col['Final Decision']])}


def is_label(value):
    return value.lower() not in UNDECIDED


def judges_agree(labels):
    return len(set(labels)) == 1 and is_label(labels[0])


def select(repos, judges, manual, label):
    """Return [(repo, source, judge_labels)] for the repos that pass. The judges decide when they
    all agree; otherwise the manual decision decides, and repos without one are left out."""
    selected = []
    for repo in repos:
        labels = [results.get(repo['url'], '') for _, results in judges]
        if judges_agree(labels):
            if labels[0] == label:
                selected.append((repo, 'judges', labels))
        elif manual.get(repo['url']) == label.lower():
            selected.append((repo, 'manual', labels))
    return selected


def write_repos_sheet(ws, judges, selected, label):
    columns = ['name', 'url', 'stars', 'language_percentage', 'frameworks',
               *(name for name, _ in judges), 'final_label', 'source']
    widths = {'name': 35, 'url': 50, 'frameworks': 45, 'language_percentage': 12}
    ws.append(columns)
    for repo, source, labels in selected:
        ws.append([repo['name'], repo['url'], repo['stars'], repo['language_percentage'],
                   ', '.join(f['name'] for f in repo.get('frameworks', [])),
                   *labels, label, source])
        ws.cell(ws.max_row, 2).hyperlink = repo['url']
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{len(selected) + 1}'
    for i, c in enumerate(columns, start=1):
        ws.cell(1, i).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 11)


def add_table(ws, row, title, header, rows, chart_title=None):
    """Write a titled table at `row`, with an optional bar chart of its 2nd column beside it.
    Returns the first free row after the section."""
    ws.cell(row, 1, title).font = Font(bold=True, size=13)
    for col, h in enumerate(header, start=1):
        cell = ws.cell(row + 1, col, h)
        cell.font, cell.fill = Font(bold=True), HEADER_FILL
    for i, values in enumerate(rows, start=row + 2):
        for col, v in enumerate(values, start=1):
            cell = ws.cell(i, col, v)
            if col > 1 and isinstance(v, float) and header[col - 1].startswith('%'):
                cell.number_format = '0.0%'
            elif isinstance(v, int):
                cell.number_format = '#,##0'
    height = len(rows) + 3
    if chart_title and rows:
        end = row + 1 + len(rows)
        chart = BarChart()
        chart.type, chart.title, chart.legend, chart.gapWidth = 'bar', chart_title, None, 60
        chart.y_axis.majorGridlines, chart.y_axis.delete, chart.x_axis.delete = None, True, False
        chart.x_axis.scaling.orientation = 'maxMin'  # keep the table order top to bottom
        chart.add_data(Reference(ws, min_col=2, min_row=row + 1, max_row=end), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=row + 2, max_row=end))
        series = chart.series[0]
        series.graphicalProperties.solidFill = '2A78D6'
        series.graphicalProperties.line.noFill = True
        series.dLbls = DataLabelList(showVal=True)
        chart.height, chart.width = max(6.0, 0.6 * len(rows) + 2), 16
        ws.add_chart(chart, f'F{row}')
        height = max(height, int(chart.height / 0.5) + 2)
    return row + height


def write_summary_sheet(ws, repos, judges, selected, manual, label):
    chosen = [repo for repo, _, _ in selected]
    total = len(chosen)
    sources = Counter(source for _, source, _ in selected)
    stars = [r['stars'] for r in chosen]

    def share(counter, limit=None):
        return [(k, n, n / total if total else 0.0) for k, n in counter.most_common(limit)]

    ws['A1'] = f'Dataset summary: repos labelled {label!r}'
    ws['A1'].font = Font(bold=True, size=16)
    ws['A2'] = f'Generated {datetime.now():%Y-%m-%d %H:%M} from {len(repos):,} input repos'
    ws['A2'].font = Font(italic=True, color='52514E')

    row = add_table(ws, 4, 'Overview', ['metric', 'value'], [
        ('Repos in dataset', total),
        ('Share of input repos', f'{total / len(repos) if repos else 0.0:.1%}'),
        ('Selected by judge agreement', sources['judges']),
        ('Selected by manual review only', sources['manual']),
    ])

    split = [r['url'] for r in repos if not judges_agree([res.get(r['url'], '') for _, res in judges])]
    row = add_table(ws, row, 'How the dataset was built', ['step', 'repos'], [
        ('Input repos', len(repos)),
        (f'All judges say {label}', sources['judges']),
        ('Judges disagree', len(split)),
        ('  of which manually reviewed', sum(u in manual for u in split)),
        (f'  of which manual review says {label}', sources['manual']),
        ('Final dataset', total),
    ])

    by_stars = Counter([name for low, name in STAR_BUCKETS if n >= low][-1] for n in stars)
    row = add_table(ws, row, 'Stars', ['stars', 'repos', '% of dataset'],
                    [(name, by_stars[name], by_stars[name] / total if total else 0.0) for _, name in STAR_BUCKETS],
                    'Repos by star count')
    row = add_table(ws, row, 'Primary language', ['language', 'repos', '% of dataset'],
                    share(Counter(r['languages'][0]['name'] if r.get('languages') else 'Unknown' for r in chosen), limit=10),
                    'Repos by primary language')
    frameworks = Counter(name for r in chosen for name in dict.fromkeys(f['name'] for f in r['frameworks']))
    row = add_table(ws, row, 'Top frameworks', ['framework', 'repos using it', '% of dataset'],
                    share(frameworks, limit=15), 'Top 15 frameworks')
    queries = Counter(r['found_via'] for r in chosen if isinstance(r.get('found_via'), str))
    row = add_table(ws, row, 'Search queries that found the most repos', ['query', 'repos', '% of dataset'],
                    share(queries, limit=10), 'Top 10 search queries')

    pairs_rows = []
    for (name_a, res_a), (name_b, res_b) in combinations(judges, 2):
        pairs = [(res_a[u], res_b[u]) for u in res_a.keys() & res_b.keys() if is_label(res_a[u]) and is_label(res_b[u])]
        agreement = sum(a == b for a, b in pairs) / len(pairs) if pairs else 0.0
        pairs_rows.append((f'{name_a} vs {name_b}', len(pairs), agreement))
    add_table(ws, row, f'Judge agreement on all {len(repos):,} input repos',
              ['judges', 'repos both labelled', '% agreement'], pairs_rows)

    ws.column_dimensions['A'].width = 48
    for c in 'BC':
        ws.column_dimensions[c].width = 18


def main():
    parser = argparse.ArgumentParser(description='Build the final dataset from judge agreement and manual review')
    parser.add_argument('judges', nargs='*',
                        help="Classification files as name=path or path (relative to the step's judges/ folder)")
    parser.add_argument('--step', choices=STEPS, default='s1', help='Screening step whose defaults to use (default: s1)')
    parser.add_argument('--label', help="Label a repo must have to pass (default: the step's, e.g. app)")
    parser.add_argument('--repos', help="Repos JSON to filter (default: the step's input, e.g. dataset/all_repos.json)")
    parser.add_argument('--manual', help="Manually reviewed spreadsheet (default: the step's manual_review.xlsx)")
    parser.add_argument('--output', '-o', help="Output path without extension; writes .xlsx and .json "
                                                "(default: the step's, e.g. dataset/s1_apps)")
    args = parser.parse_args()

    step = STEPS[args.step]
    step_dir = os.path.join(SCREENING_DIR, step['dir'])
    judges_dir = os.path.join(step_dir, 'judges')
    args.label = args.label or step['label']
    args.repos = args.repos or os.path.join(DATASET_DIR, step['repos'])
    args.manual = args.manual or os.path.join(step_dir, 'manual_review.xlsx')
    args.output = args.output or os.path.join(DATASET_DIR, step['output'])
    specs = args.judges or step['judges'] or [
        f'{os.path.basename(p)[:-5]}={p}' for p in sorted(glob.glob(os.path.join(judges_dir, '*.json')))]
    if not specs:
        parser.error(f'no judge files given or found in {judges_dir}')

    judges = [load_judge(spec, judges_dir) for spec in specs]
    with open(args.repos) as f:
        all_repos = json.load(f)
    repos = all_repos['repos']
    manual = load_manual(resolve_path(args.manual, step_dir))
    selected = select(repos, judges, manual, args.label)
    sources = Counter(source for _, source, _ in selected)

    wb = Workbook()
    wb.active.title = 'Repos'
    write_repos_sheet(wb.active, judges, selected, args.label)
    write_summary_sheet(wb.create_sheet('Summary', 0), repos, judges, selected, manual, args.label)
    wb.save(args.output + '.xlsx')

    metadata = {**all_repos['metadata'], 'total_repos': len(selected), 'last_updated': datetime.now().isoformat(),
                'label': args.label, 'judges': [name for name, _ in judges],
                'selected_by_judges': sources['judges'], 'selected_by_manual_review': sources['manual']}
    with open(args.output + '.json', 'w') as f:
        json.dump({'metadata': metadata, 'repos': [repo for repo, _, _ in selected]}, f, indent=2, ensure_ascii=False)

    print(f'{len(selected)} of {len(repos)} repos labelled {args.label!r} '
          f'(judges agree: {sources["judges"]}, manual review: {sources["manual"]})')
    split = [r['url'] for r in repos if not judges_agree([res.get(r['url'], '') for _, res in judges])]
    print(f'  judges disagree on {len(split)} repos, {sum(u in manual for u in split)} of them manually reviewed')
    print(f'Wrote {args.output}.xlsx and {args.output}.json')


if __name__ == '__main__':
    main()
